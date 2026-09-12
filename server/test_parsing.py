"""解析器测试：测试内现场构造合成 xlsx/docx，不依赖真实样例文件。"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from parsing import ParseError, parse_docx, parse_file, parse_weeks, rows_to_entries  # noqa: E402


def build_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_docx(rows: list[list[str]]) -> bytes:
    from docx import Document

    document = Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for col_index, cell in enumerate(row):
            table.rows[row_index].cells[col_index].text = cell
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


STANDARD_TABLE = [
    ["课程名称", "教师", "星期", "周次", "节次", "地点"],
    ["生态毒理学", "王咏", "周一", "1-16周", "0102", "环境学院216室"],
    ["保护生物学", "王平", "星期三", "1,3,5周", "1-2节", "环境学院215室"],
    ["水文与水资源", "杨武", "周五", "2-16双周", "第3,4节", "环境学院216室"],
]


def test_standard_xlsx_headers():
    entries = parse_file("课表.xlsx", build_xlsx(STANDARD_TABLE))

    first = entries[0]
    assert first["course"] == "生态毒理学"
    assert first["teacher"] == "王咏"
    assert first["weekday"] == 1
    assert first["weekday_name"] == "周一"
    assert first["periods"] == [1, 2]
    assert first["period_code"] == "0102"
    assert first["location"] == "环境学院216室"
    assert first["course_type"] == "理论课时"
    # 1-16周 展开出 16 条
    course_entries = [e for e in entries if e["course"] == "生态毒理学"]
    assert len(course_entries) == 16
    assert [e["week"] for e in course_entries] == list(range(1, 17))
    # date 由学期起点推算：第 1 周 = 2026-09-07
    assert course_entries[0]["date"] == "2026-09-07"
    assert course_entries[1]["date"] == "2026-09-14"
    assert course_entries[0]["big_sections"] == [1]


def test_week_variant_formats():
    assert parse_weeks("1-16周") == list(range(1, 17))
    assert parse_weeks("1,3,5周") == [1, 3, 5]
    assert parse_weeks("1-15单周") == [1, 3, 5, 7, 9, 11, 13, 15]
    assert parse_weeks("2-16双周") == [2, 4, 6, 8, 10, 12, 14, 16]
    assert parse_weeks("第3周") == [3]


def test_docx_table_path():
    entries = parse_file("课表.docx", build_docx(STANDARD_TABLE))
    assert entries[0]["course"] == "生态毒理学"
    assert any(e["course"] == "水文与水资源" and e["weekday"] == 5 for e in entries)


def test_unsupported_extension():
    with pytest.raises(ParseError, match="暂不支持"):
        parse_file("课表.png", b"fake")


def test_no_header_raises():
    with pytest.raises(ParseError):
        rows_to_entries([["随便", "什么"], ["都", "没有"]])


def test_empty_rows_dropped():
    rows = [
        ["课程名称", "教师", "星期", "周次", "节次", "地点"],
        ["", "", "", "", "", ""],
        ["海洋生态学", "李某", "周二", "1-8周", "0304", "教学楼"],
    ]
    entries = rows_to_entries(rows)
    assert len(entries) == 8
    assert all(e["course"] == "海洋生态学" for e in entries)
